from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from PIL import UnidentifiedImageError


_SNAPSHOT_PATTERNS = {
    "texture": "step*_texture_operation.png",
    "gaussian": "step*_gaussian_rendering.png",
    "ai_retexture": "step*_ai_retexture_operation.png",
    "human_body_completion": "step*_human_body_completion.png",
}
_IMAGE_MAX_WIDTH = 500
_IMAGE_MAX_HEIGHT = 282
_FONT_NAME = "Arial"

# 各后处理类型对应的 CrealityScan 独立下载包目录名（Extensions 下的子目录）。
# 高斯渲染与 AI重贴图共用同一个下载包；贴图无独立下载包。
# 头模置换（head_restore）对应 AIHeadRestore，工具尚未纳入该后处理类型，暂不配置。
_DOWNLOAD_PACKAGE_DIRS = {
    "texture": None,
    "gaussian": "AITextureRestore",
    "ai_retexture": "AITextureRestore",
    "human_body_completion": "AIBodyComplete",
}
_DOWNLOAD_PACKAGE_LABELS = {
    "texture": None,
    "gaussian": "高斯渲染下载包",
    "ai_retexture": "AI重贴图下载包",
    "human_body_completion": "人体补全下载包",
}


@dataclass(frozen=True)
class ReportProject:
    sequence: int
    source_project_dir: Path
    release_run_dir: Path
    test_run_dir: Path


def _find_snapshot(run_dir: Path, operation: str) -> tuple[Optional[Path], str]:
    pattern = _SNAPSHOT_PATTERNS.get(operation)
    if pattern is None:
        raise ValueError(f"不支持的后处理类型：{operation}")
    matches = sorted((run_dir / "screenshots").glob(pattern))
    if not matches:
        return None, "未找到截图"
    if len(matches) > 1:
        return None, f"截图数量异常（{len(matches)} 张）"
    return matches[0], "已嵌入"


def _add_snapshot(worksheet: Worksheet, cell: str, snapshot_path: Path) -> str:
    try:
        image = ExcelImage(str(snapshot_path))
    except (OSError, ValueError, UnidentifiedImageError):
        worksheet[cell] = "截图无法读取"
        return "截图无法读取"

    scale = min(
        1.0,
        _IMAGE_MAX_WIDTH / float(image.width),
        _IMAGE_MAX_HEIGHT / float(image.height),
    )
    image.width = int(image.width * scale)
    image.height = int(image.height * scale)
    worksheet.add_image(image, cell)
    return "已嵌入"


def _style_report(worksheet: Worksheet, project_count: int) -> None:
    navy = "18324A"
    blue = "2F6FED"
    pale_blue = "EAF1FF"
    pale_gray = "F4F6F8"
    border_color = "CBD5E1"
    thin = Side(style="thin", color=border_color)

    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 75
    worksheet.freeze_panes = "B4"
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 72
    worksheet.column_dimensions["C"].width = 72
    worksheet.column_dimensions["D"].width = 20
    worksheet.row_dimensions[1].height = 32
    worksheet.row_dimensions[2].height = 25
    worksheet.row_dimensions[3].height = 26

    title = worksheet["A1"]
    title.font = Font(name=_FONT_NAME, size=18, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    for row in worksheet["A1:D1"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=navy)

    metadata = worksheet["A2"]
    metadata.font = Font(name=_FONT_NAME, size=10, color="334155")
    metadata.fill = PatternFill("solid", fgColor=pale_blue)
    metadata.alignment = Alignment(horizontal="left", vertical="center")
    for row in worksheet["A2:D2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=pale_blue)

    for cell in worksheet[3]:
        cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_number in range(4, 4 + project_count):
        worksheet.row_dimensions[row_number].height = (_IMAGE_MAX_HEIGHT * 0.75) + 14
        for column_number in range(1, 5):
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.font = Font(name=_FONT_NAME, size=9, color="1E293B")
            cell.fill = PatternFill(
                "solid",
                fgColor="FFFFFF" if row_number % 2 == 0 else pale_gray,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    worksheet.auto_filter.ref = f"A3:D{max(3, project_count + 3)}"
    worksheet.print_options.horizontalCentered = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_area = f"A1:D{max(3, project_count + 3)}"


def _extensions_root() -> Path:
    """返回 CrealityScan Extensions 下载包根目录。

    读取目录：%LOCALAPPDATA%\\Creality\\CrealityScan\\Extensions
    """
    local_appdata = os.environ.get("LOCALAPPDATA")
    base = Path(local_appdata) if local_appdata else Path.home() / "AppData" / "Local"
    return base / "Creality" / "CrealityScan" / "Extensions"


def read_download_package_version(operation: str) -> Optional[str]:
    """读取 CrealityScan 下载包 version.txt 中的版本号。

    无独立下载包的后处理类型（如贴图）返回 None；读取失败也返回 None。
    """
    package_dir = _DOWNLOAD_PACKAGE_DIRS.get(operation)
    if not package_dir:
        return None
    version_file = _extensions_root() / package_dir / "version.txt"
    try:
        version = version_file.read_text(encoding="utf-8", errors="replace").strip()
    except OSError:
        return None
    return version or None


def iter_download_package_dirs(operation: str) -> list[Path]:
    """返回该后处理操作对应的 CrealityScan 下载包目录（含 _Data 数据目录）绝对路径列表。

    无独立下载包的后处理类型（如贴图）返回空列表。用于发布版跑完后删除旧包，
    使测试版运行前重新触发下载新版本下载包。
    """
    package_dir = _DOWNLOAD_PACKAGE_DIRS.get(operation)
    if not package_dir:
        return []
    root = _extensions_root()
    return [root / package_dir, root / f"{package_dir}_Data"]


def generate_excel_report(
    run_root: Path,
    operation: str,
    operation_name: str,
    release_version: str,
    test_version: str,
    projects: Sequence[ReportProject],
    generated_at: Optional[datetime] = None,
) -> Path:
    if not projects:
        raise ValueError("没有可写入 Excel 对比表的工程。")
    if operation not in _SNAPSHOT_PATTERNS:
        raise ValueError(f"不支持的后处理类型：{operation}")

    report_time = generated_at or datetime.now()
    report_path = run_root / f"{report_time.strftime('%Y%m%d_%H%M%S')}{operation_name}对比表.xlsx"
    temporary_path = report_path.with_name(f".{report_path.stem}.tmp.xlsx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "截图对比"
    worksheet.merge_cells("A1:D1")
    worksheet.merge_cells("A2:D2")
    worksheet["A1"] = f"{operation_name}截图对比表"
    package_label = _DOWNLOAD_PACKAGE_LABELS.get(operation)
    package_note = ""
    if package_label:
        package_version = read_download_package_version(operation)
        package_note = f"    {package_label}：{package_version or '未获取'}"
    worksheet["A2"] = (
        f"生成时间：{report_time.strftime('%Y-%m-%d %H:%M:%S')}    "
        f"发布版：{release_version}    测试版：{test_version}    "
        f"工程数：{len(projects)}"
        f"{package_note}"
    )
    worksheet.append(
        [
            "工程",
            f"发布版截图\n{release_version}",
            f"测试版截图\n{test_version}",
            "截图状态",
        ]
    )

    for row_number, project in enumerate(projects, start=4):
        worksheet.cell(
            row=row_number,
            column=1,
            value=(
                f"工程 {project.sequence}\n"
                f"{project.source_project_dir.name}\n"
                f"{project.source_project_dir}"
            ),
        )
        release_path, release_status = _find_snapshot(project.release_run_dir, operation)
        test_path, test_status = _find_snapshot(project.test_run_dir, operation)
        if release_path is not None:
            release_status = _add_snapshot(worksheet, f"B{row_number}", release_path)
        else:
            worksheet.cell(row=row_number, column=2, value=release_status)
        if test_path is not None:
            test_status = _add_snapshot(worksheet, f"C{row_number}", test_path)
        else:
            worksheet.cell(row=row_number, column=3, value=test_status)

        status = "截图完整"
        if release_status != "已嵌入" or test_status != "已嵌入":
            status = f"发布版：{release_status}\n测试版：{test_status}"
        worksheet.cell(row=row_number, column=4, value=status)

    _style_report(worksheet, len(projects))
    workbook.properties.title = f"{operation_name}截图对比表"
    workbook.properties.subject = "CrealityScan 发布版与测试版截图对比"
    workbook.properties.creator = "CrealityScan 后处理对比平台"

    run_root.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(report_path)
    except (OSError, TypeError, ValueError) as exc:
        if temporary_path.exists():
            temporary_path.unlink()
        raise RuntimeError(f"Excel 对比表保存失败：{exc}") from exc
    return report_path
