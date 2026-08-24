from __future__ import annotations

from pathlib import Path
from typing import Any, Optional


WORKBOOK_NAME = "scan_fps_summary.xlsx"
SHEET_TITLE = "扫描帧率"
TABLE_NAME = "ScanFpsTable"
HEADERS = [
    "序号",
    "任务名称",
    "扫描模组",
    "连接方式",
    "参数配置组",
    "参数配置步骤",
    "扫描步骤",
    "开始时间",
    "结束时间",
    "平均帧率(FPS)",
    "融合耗时(秒)",
    "封装耗时(秒)",
    "贴图耗时(秒)",
    "目标帧数",
    "起始帧",
    "停止帧",
    "帧增量",
    "扫描耗时(秒)",
    "结果状态",
    "失败原因",
]


def _number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str:
    return str(value or "").strip()


def _nearest_scan_params(case_steps: Any, result_index: int) -> tuple[str, str]:
    if not isinstance(case_steps, list):
        return "", ""
    for candidate in reversed(case_steps[:result_index]):
        if not isinstance(candidate, dict):
            continue
        step_id = _text(candidate.get("id"))
        if not step_id.startswith("crealityscan.configure_scan_params"):
            continue
        candidate_params = candidate.get("params")
        candidate_params = candidate_params if isinstance(candidate_params, dict) else {}
        preset = _text(candidate_params.get("preset"))
        return preset or _text(candidate.get("name")) or step_id, step_id
    return "", ""


def _module_name(config_step_id: str) -> str:
    marker = "crealityscan.configure_scan_params_"
    raw_name = config_step_id[len(marker) :].strip() if config_step_id.startswith(marker) else ""
    display_names = {
        "feeret": "Feeret",
        "otter": "Otter",
        "otter_lite": "Otter Lite",
        "otter_lite_basic": "Otter Lite Basic",
        "p1": "P1",
        "p1s": "P1S",
        "pika": "Pika",
        "raptor": "Raptor",
        "raptor_x": "Raptor X",
        "raptor_pro": "Raptor Pro",
        "s1": "S1",
        "x1": "X1",
    }
    return display_names.get(raw_name.lower(), raw_name or "")


def _connection_type(case: dict[str, Any], device_info: Optional[dict[str, Any]]) -> str:
    info = device_info if isinstance(device_info, dict) else {}
    app = case.get("app") if isinstance(case.get("app"), dict) else {}
    return (
        _text(info.get("camera_connection_type"))
        or _text(app.get("connection_type"))
        or _text(app.get("device_uri"))
    )


def _failure_reason(item: dict[str, Any]) -> str:
    error = _text(item.get("error"))
    if not error:
        return ""
    return error.splitlines()[-1][:500]


def _post_processing_durations(step_results: list[dict[str, Any]], result_index: int) -> dict[str, Optional[float]]:
    operation_ids = {
        "crealityscan.fusion_operation": "fusion",
        "crealityscan.package_operation": "package",
        "crealityscan.texture_operation": "texture",
    }
    durations: dict[str, Optional[float]] = {key: None for key in operation_ids.values()}
    for candidate in step_results[result_index + 1 :]:
        if not isinstance(candidate, dict):
            continue
        candidate_id = _text(candidate.get("id"))
        if (
            candidate_id in {"crealityscan.create_scan", "crealityscan.create_project"}
            or candidate_id == "crealityscan.scan_until_frames_then_stop"
            or candidate_id.startswith("crealityscan.configure_scan_params")
        ):
            break
        operation_key = operation_ids.get(candidate_id)
        if not operation_key:
            continue
        candidate_extra = candidate.get("extra") if isinstance(candidate.get("extra"), dict) else {}
        duration = _number(candidate_extra.get("operation_elapsed_sec"))
        if duration is None:
            duration = _number(candidate.get("duration_sec"))
        durations[operation_key] = duration
    return durations


def _build_rows(
    case: dict[str, Any],
    step_results: list[dict[str, Any]],
    device_info: Optional[dict[str, Any]],
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    case_steps = case.get("steps")
    case_name = _text(case.get("name") or case.get("case_id"))
    scan_number = 0

    for result_index, item in enumerate(step_results):
        if not isinstance(item, dict):
            continue
        if _text(item.get("id")) != "crealityscan.scan_until_frames_then_stop":
            continue
        status = _text(item.get("status"))
        if status not in {"passed", "failed"}:
            continue

        scan_number += 1
        extra = item.get("extra") if isinstance(item.get("extra"), dict) else {}
        preset, preset_step_id = _nearest_scan_params(case_steps, result_index)
        post_processing = _post_processing_durations(step_results, result_index)
        rows.append(
            [
                scan_number,
                case_name,
                _module_name(preset_step_id),
                _connection_type(case, device_info),
                preset,
                preset_step_id,
                _text(item.get("name") or item.get("id")),
                _text(extra.get("scan_started_log_at")),
                _text(extra.get("sdk_fps_session_stop_at") or extra.get("stop_click_log_at")),
                _number(extra.get("avg_fps")),
                post_processing["fusion"],
                post_processing["package"],
                post_processing["texture"],
                _number(extra.get("target_frames")),
                _number(extra.get("start_frame")),
                _number(extra.get("stop_click_frame")),
                _number(extra.get("frame_delta")),
                _number(extra.get("scan_elapsed_sec")),
                status,
                _failure_reason(item),
            ]
        )
    return rows


def _style_worksheet(ws: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(HEADERS)):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for row_index in range(2, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 20

    for row_index in range(2, ws.max_row + 1):
        ws.cell(row_index, 8).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row_index, 9).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in (10, 11, 12, 13, 18):
            ws.cell(row_index, column).number_format = "0.000"
        for column in (14, 15, 16, 17):
            ws.cell(row_index, column).number_format = "0"

    widths = {
        "A": 6,
        "B": 16,
        "C": 10,
        "D": 12,
        "E": 24,
        "F": 31,
        "G": 18,
        "H": 20,
        "I": 20,
        "J": 13,
        "K": 12,
        "L": 12,
        "M": 12,
        "N": 10,
        "O": 10,
        "P": 10,
        "Q": 10,
        "R": 12,
        "S": 10,
        "T": 32,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def write_scan_fps_workbook(
    case: dict[str, Any],
    step_results: list[dict[str, Any]],
    run_dir: str | Path,
    device_info: Optional[dict[str, Any]] = None,
) -> Optional[str]:
    """Write the SDK FPS already attached to scan step results into the run artifact."""
    rows = _build_rows(case, step_results, device_info)
    if not rows:
        return None

    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("生成扫描帧率 Excel 需要 openpyxl，请安装项目依赖。") from exc

    run_path = Path(run_dir)
    run_path.mkdir(parents=True, exist_ok=True)
    workbook_path = run_path / WORKBOOK_NAME
    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
        worksheet = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.create_sheet(SHEET_TITLE)
        if worksheet.max_row == 1 and worksheet.cell(1, 1).value is None:
            worksheet.delete_rows(1, 1)
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = SHEET_TITLE

    if worksheet.max_row == 0 or worksheet.cell(1, 1).value is None:
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(1, column).value = header
        worksheet.cell(1, 3).comment = Comment(
            "来源：本次扫描前最近一次 crealityscan.configure_scan_params_* 步骤 ID 的模组后缀。",
            "Jens",
        )
        worksheet.cell(1, 4).comment = Comment(
            "来源：运行产物 device_info.camera_connection_type；缺失时回退到任务连接配置。",
            "Jens",
        )
        worksheet.cell(1, 5).comment = Comment(
            "来源：本次扫描前最近一次 crealityscan.configure_scan_params_* 步骤的 preset。",
            "Jens",
        )
        worksheet.cell(1, 10).comment = Comment(
            "来源：步骤结果 extra.avg_fps；该值由 Runner 已完成的 SDK FPS 会话匹配提供。",
            "Jens",
        )
        worksheet.cell(1, 11).comment = Comment(
            "来源：后处理步骤 extra.operation_elapsed_sec，步骤自身的融合操作耗时。",
            "Jens",
        )
        worksheet.cell(1, 12).comment = Comment(
            "来源：后处理步骤 extra.operation_elapsed_sec，步骤自身的封装操作耗时。",
            "Jens",
        )
        worksheet.cell(1, 13).comment = Comment(
            "来源：后处理步骤 extra.operation_elapsed_sec，步骤自身的贴图操作耗时。",
            "Jens",
        )
        worksheet.cell(1, 19).comment = Comment(
            "passed/failed 表示扫描步骤执行结果；skipped 步骤不会写入。",
            "Jens",
        )
    elif [worksheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)] != HEADERS:
        raise ValueError(f"{workbook_path} 的工作表列结构不兼容，无法追加扫描帧率结果。")

    for row in rows:
        worksheet.append(row)

    table = next((item for item in worksheet.tables.values() if item.name == TABLE_NAME), None)
    table_ref = f"A1:T{worksheet.max_row}"
    if table is None:
        table = Table(displayName=TABLE_NAME, ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        table.ref = table_ref

    _style_worksheet(worksheet)
    temp_path = workbook_path.with_suffix(".tmp.xlsx")
    workbook.save(temp_path)
    temp_path.replace(workbook_path)
    return str(workbook_path)
