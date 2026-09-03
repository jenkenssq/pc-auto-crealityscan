from __future__ import annotations

import math
import re
import shutil
from pathlib import Path
from typing import Any, Optional, Sequence

WORKBOOK_NAME = "帧率统计.xlsx"
SHEET_TITLE = "Sheet1"
FONT_NAME = "Microsoft YaHei"
FONT_SIZE = 14

# 与 帧率统计模板.xlsx 一致的表头与行列尺寸。
HEADERS = ["软件版本", "电脑", "系统版本", "wifi5/6", "电脑型号", "内存", "", ""]
ROW_HEIGHTS = {1: 83.0, 2: 146.0, 3: 266.5}
COLUMN_WIDTHS = {
    "A": 43.66,
    "B": 21.08,
    "C": 19.83,
    "D": 13.58,
    "E": 32.08,
    "F": 8.5,
    "G": 31.66,
    "H": 34.33,
}


def truncate_int(value: Any) -> int:
    """保留整数部分（直接去掉小数），不做四舍五入。"""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0
    if math.isnan(number):
        return 0
    return int(number)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _template_has_sample_data(template_path: Path) -> bool:
    """判断模板是否带示例数据（如 帧率统计模板.xlsx 的 G3/H3 含示例帧率行）。"""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(template_path, data_only=True)
        ws = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.active
        for coordinate in ("G3", "H3"):
            if _text(ws[coordinate].value):
                return True
        return False
    except Exception:
        return True


def _normalize_connection(connection_type: str) -> str:
    token = (connection_type or "").strip().lower().replace("_", "-")
    if token in {"wifi", "wi-fi", "wlan"}:
        return "WIFI"
    if token in {"usb"}:
        return "USB"
    return ""


def _pc_short_name(system_info: Optional[dict]) -> str:
    """电脑简称：从 GPU 型号取数字，如 RTX 5060 → 5060笔记本。"""
    info = system_info or {}
    gpu = _text(info.get("gpu"))
    match = re.search(r"(\d{3,4})", gpu)
    if match:
        return f"{match.group(1)}笔记本"
    cpu = _text(info.get("cpu"))
    match = re.search(r"(\d{3,4})", cpu)
    if match:
        return f"{match.group(1)}笔记本"
    return "笔记本"


def _system_version_text(system_info: Optional[dict]) -> str:
    """系统版本：Windows 11 专业版 \\n 23H2 \\n 22631（Q4 自动补全专业版/版本名）。"""
    info = system_info or {}
    lines = [
        _text(info.get("os_name")),
        _text(info.get("os_version_name")),
        _text(info.get("os_build")),
    ]
    lines = [line for line in lines if line]
    return "\n".join(lines) or _text(info.get("os_raw"))


def _pc_model_text(system_info: Optional[dict]) -> str:
    info = system_info or {}
    return "\n".join(
        [
            f"CPU: {info.get('cpu', '')}",
            f"CPU Cores: {info.get('cpu_cores', '')}",
            f"Memory: {info.get('memory_mb', '')}MB",
            f"GPU: {info.get('gpu', '')}",
            f"GPU Memory: {info.get('gpu_memory', '')}",
        ]
    )


def _memory_text(system_info: Optional[dict]) -> str:
    """内存：按 GB 向上取整，与模板（32175MB → 32g）一致。"""
    info = system_info or {}
    try:
        memory_mb = int(info.get("memory_mb") or 0)
    except (TypeError, ValueError):
        memory_mb = 0
    if memory_mb <= 0:
        return ""
    return f"{math.ceil(memory_mb / 1024)}g"


def _format_fps_line(
    mode_name: str,
    preview: Any,
    scan_min: Any,
    scan_max: Any,
    stable: Any,
) -> Optional[str]:
    """拼一行“模式名：预览；最小-最大；稳定”。全部缺值返回 None（跳过该行）。"""
    cells = [preview, scan_min, scan_max, stable]
    if all(value is None or value == "" for value in cells):
        return None

    def _cell(value: Any) -> str:
        return "" if value is None or value == "" else str(truncate_int(value))

    return f"{mode_name}：{_cell(preview)}；{_cell(scan_min)}-{_cell(scan_max)}；{_cell(stable)}"


def build_fps_rows(mode_names: Sequence[str], phase_blocks: Sequence[dict]) -> list[dict]:
    """把分阶段提取结果按模式顺序映射为 fps_rows（模式名 + 四个数值）。"""
    rows: list[dict] = []
    for index, name in enumerate(mode_names):
        block = phase_blocks[index] if index < len(phase_blocks) else {}
        block = block if isinstance(block, dict) else {}
        rows.append(
            {
                "mode_name": name,
                "preview": block.get("preview_fps"),
                "scan_min": block.get("scan_min_fps"),
                "scan_max": block.get("scan_max_fps"),
                "stable": block.get("scan_avg_fps"),
            }
        )
    return rows


def _write_headers(ws: Any) -> None:
    for column, header in enumerate(HEADERS, start=1):
        ws.cell(1, column).value = header or None


def _style_workbook(ws: Any) -> None:
    from openpyxl.styles import Alignment, Font

    header_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    body_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=False)
    for row_index in (1, 2, 3):
        ws.row_dimensions[row_index].height = ROW_HEIGHTS.get(row_index, 20)
    for column in COLUMN_WIDTHS:
        ws.column_dimensions[column].width = COLUMN_WIDTHS[column]

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A2"].font = header_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for coordinate in ("B3", "C3", "D3", "E3", "F3", "G3", "H3"):
        cell = ws[coordinate]
        cell.font = body_font
        horizontal = "left" if coordinate in {"G3", "H3"} else "center"
        cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


def write_fps_stat_workbook(
    output_path: str | Path,
    *,
    software_version: str,
    module_display_name: str,
    connection_type: str,
    system_info: Optional[dict] = None,
    firmware_version: str = "",
    wifi_handle_version: str = "",
    wifi_band: str = "",
    fps_rows: Optional[Sequence[dict]] = None,
    template_path: Optional[str | Path] = None,
) -> str:
    """按 帧率统计模板.xlsx 写入帧率统计结果。

    - WIFI → 写 G 列（含 G1 表头），USB → 写 H 列（含 H1 表头）；
    - 已存在的工作簿只更新系统信息与目标列，保留另一列已有数据；
    - 无模板时按模板样式自建空工作簿。
    """
    try:
        from openpyxl import load_workbook, Workbook
    except ImportError as exc:
        raise RuntimeError("生成帧率统计 Excel 需要 openpyxl，请安装项目依赖。") from exc

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    connection = _normalize_connection(connection_type)
    if connection not in {"WIFI", "USB"}:
        raise ValueError(f"无法识别的连接方式：{connection_type}（需为 Wi-Fi 或 USB）")

    # 只有模板是“干净模板”（G/H 列无示例数据）时才复用其样式结构，
    # 避免把 帧率统计模板.xlsx 里的示例数值带进真实产物。
    if (
        not target.exists()
        and template_path
        and Path(template_path).is_file()
        and not _template_has_sample_data(Path(template_path))
    ):
        shutil.copyfile(Path(template_path), target)

    if target.exists():
        workbook = load_workbook(target)
        ws = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.create_sheet(SHEET_TITLE)
        if ws.max_row < 3 or ws.cell(1, 1).value is None:
            _write_headers(ws)
    else:
        workbook = Workbook()
        ws = workbook.active
        ws.title = SHEET_TITLE
        _write_headers(ws)

    # 系统信息（公共区）
    try:
        ws.merge_cells("A2:A3")
    except ValueError:
        pass
    ws["A2"] = f"win{_text(software_version)}"
    system_info = system_info or {}
    ws["B3"] = _pc_short_name(system_info)
    ws["C3"] = _system_version_text(system_info)
    # wifi5/6 只在 WIFI 连接时填写（USB 留空）
    ws["D3"] = _text(wifi_band) if connection == "WIFI" else ""
    ws["E3"] = _pc_model_text(system_info)
    ws["F3"] = _memory_text(system_info)

    # 目标列表头与内容（只写目标列，另一列保留不动）
    target_column = "G" if connection == "WIFI" else "H"
    if connection == "WIFI":
        ws[f"{target_column}1"] = (
            f"{_text(module_display_name)} WIFI（固件版本：{_text(firmware_version)}+wifi ：{_text(wifi_handle_version)}）"
        )
    else:
        ws[f"{target_column}1"] = f"{_text(module_display_name)}(USB) 固件版本：{_text(firmware_version)}"

    if fps_rows:
        lines = [_format_fps_line(**row) for row in fps_rows]
        ws[f"{target_column}3"] = "\n".join(line for line in lines if line)
    else:
        ws[f"{target_column}3"] = ""

    _style_workbook(ws)
    temp_path = target.with_suffix(target.suffix + ".tmp")
    workbook.save(temp_path)
    temp_path.replace(target)
    return str(target)
