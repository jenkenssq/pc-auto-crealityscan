import tempfile
import unittest
from pathlib import Path

from engine.fps_stat_xlsx import (
    build_fps_rows,
    truncate_int,
    write_fps_stat_workbook,
)
from engine.logs import extract_fps_metrics_by_phase
from engine.system_env import (
    parse_system_info_text,
    read_software_version,
    read_wifi_handle_version,
    read_system_info_from_logs,
)
from jens_platform.task_generator import (
    TASK_KIND_FPS_STAT,
    build_task_model,
    fps_stat_mode_plan,
    fps_stat_preset_keys,
)

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
_SAMPLE_ROOT = _PROJECT_ROOT / "帧率统计任务生成方案" / "20260901112214"


class FpsStatPhaseExtractionTests(unittest.TestCase):
    def test_phase_metrics_match_sample(self) -> None:
        blocks = extract_fps_metrics_by_phase(str(_SAMPLE_ROOT))
        self.assertEqual(len(blocks), 1)
        block = blocks[0]
        self.assertAlmostEqual(block["preview_fps"], 40.621269, places=3)
        self.assertAlmostEqual(block["scan_min_fps"], 38.942978, places=3)
        self.assertAlmostEqual(block["scan_max_fps"], 40.751099, places=3)
        self.assertAlmostEqual(block["scan_avg_fps"], 39.95864157142857, places=3)
        self.assertEqual(block["preview_samples"], 1)
        self.assertEqual(block["scan_samples"], 7)

    def test_truncate_int(self) -> None:
        self.assertEqual(truncate_int(40.62), 40)
        self.assertEqual(truncate_int(38.94), 38)
        self.assertEqual(truncate_int(39.96), 39)
        self.assertEqual(truncate_int(39.5), 39)


class FpsStatModePlanTests(unittest.TestCase):
    def test_raptor_pro_wifi_excludes_no_marker_and_single(self) -> None:
        plan = fps_stat_mode_plan("raptor pro", "Wi-Fi")
        names = [name for name, _key in plan]
        self.assertEqual(names, ["平行线", "交叉", "大物体", "中物体", "小物体", "人脸", "人体"])

    def test_raptor_pro_usb_includes_no_marker(self) -> None:
        names = [name for name, _key in fps_stat_mode_plan("raptor pro", "USB")]
        self.assertIn("无标记点", names)

    def test_p1_usb_has_laser_lines_and_geometry_sizes(self) -> None:
        plan = fps_stat_mode_plan("P1", "USB")
        names = [name for name, _key in plan]
        self.assertEqual(
            names[:5],
            ["平行线", "单线", "交叉", "无标记点-交叉线", "无标记点-平行线"],
        )
        self.assertIn("p1.line_laser.no_marker.cross", [key for _, key in plan])
        self.assertIn("p1.line_laser.no_marker.parallel", [key for _, key in plan])
        for size in ("大物体", "中物体", "小物体"):
            self.assertIn(size, names)

    def test_pika_usb_includes_with_marker_line_laser_modes(self) -> None:
        plan = fps_stat_mode_plan("pika", "USB")
        names = [name for name, _key in plan]
        self.assertEqual(
            names[:4],
            ["有标志点-标准", "有标志点-均衡", "有标志点-快速", "无标记点"],
        )
        self.assertIn("pika.line_laser.point_cloud.with_marker.standard", [key for _, key in plan])
        self.assertIn("pika.line_laser.point_cloud.with_marker.balanced", [key for _, key in plan])
        self.assertIn("pika.line_laser.point_cloud.with_marker.fast", [key for _, key in plan])
        self.assertIn("pika.line_laser.point_cloud.no_marker", [key for _, key in plan])

    def test_pika_wifi_excludes_with_marker_and_no_marker_line_laser(self) -> None:
        keys = [key for _name, key in fps_stat_mode_plan("pika", "Wi-Fi")]
        self.assertTrue(all("line_laser" not in key for key in keys))

    def test_fps_stat_build_task_model_uses_1000_frames_and_persists_metadata(self) -> None:
        model = build_task_model(
            "raptor pro",
            TASK_KIND_FPS_STAT,
            "Raptor Pro帧率统计wifi",
            connection_type="Wi-Fi",
            target_frames=1000,
        )
        data = model.to_dict()
        self.assertEqual(data["task_kind"], TASK_KIND_FPS_STAT)
        self.assertEqual(data["connection_type"], "Wi-Fi")
        frames = [s["params"].get("target_frames") for s in data["steps"] if s["id"].endswith("scan_until_frames_then_stop")]
        self.assertTrue(frames)
        self.assertTrue(all(frame == 1000 for frame in frames))
        # WIFI 下不应包含 no_marker preset
        keys = [s["params"].get("preset") for s in data["steps"] if s["id"].startswith("crealityscan.configure_scan_params")]
        self.assertTrue(all("no_marker" not in (key or "") for key in keys))


class FpsStatWorkbookTests(unittest.TestCase):
    def test_wifi_writes_g_column_and_leaves_h_empty(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "帧率统计.xlsx"
            write_fps_stat_workbook(
                output,
                software_version="1.12.11",
                module_display_name="Creality Raptor Pro",
                connection_type="Wi-Fi",
                system_info={
                    "cpu": "CPU X",
                    "cpu_cores": "8",
                    "memory_mb": 32175,
                    "gpu": "NVIDIA GeForce RTX 5060 Laptop GPU",
                    "gpu_memory": "7899MB",
                    "os_name": "Windows 11 专业版",
                    "os_version_name": "23H2",
                    "os_build": "22631",
                },
                firmware_version="1.4.9",
                wifi_handle_version="1.3.1",
                wifi_band="wifi6",
                fps_rows=[{"mode_name": "平行线", "preview": 40.62, "scan_min": 38.94, "scan_max": 40.75, "stable": 39.96}],
            )
            from openpyxl import load_workbook

            ws = load_workbook(output)["Sheet1"]
            self.assertEqual(ws["A2"].value, "win1.12.11")
            self.assertEqual(ws["F3"].value, "32g")
            self.assertEqual(
                ws["G1"].value,
                "Creality Raptor Pro WIFI（固件版本：1.4.9+wifi ：1.3.1）",
            )
            self.assertEqual(ws["G3"].value, "平行线：40；38-40；39")
            self.assertIsNone(ws["H1"].value)
            self.assertIsNone(ws["H3"].value)

    def test_accumulate_second_connection_preserves_first_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "帧率统计.xlsx"
            kwargs = dict(
                software_version="1.12.11",
                module_display_name="Creality Raptor Pro",
                system_info={},
                firmware_version="1.4.9",
                wifi_handle_version="1.3.1",
                wifi_band="",
                fps_rows=[{"mode_name": "平行线", "preview": 40.62, "scan_min": 38.94, "scan_max": 40.75, "stable": 39.96}],
            )
            write_fps_stat_workbook(output, connection_type="Wi-Fi", **kwargs)
            write_fps_stat_workbook(output, connection_type="USB", **kwargs)
            from openpyxl import load_workbook

            ws = load_workbook(output)["Sheet1"]
            self.assertEqual(ws["G3"].value, "平行线：40；38-40；39")
            self.assertEqual(ws["H3"].value, "平行线：40；38-40；39")
            self.assertIn("(USB) 固件版本：1.4.9", ws["H1"].value)

    def test_empty_mode_rows_are_skipped(self) -> None:
        rows = build_fps_rows(["平行线", "交叉"], [{"preview_fps": 1.0, "scan_min_fps": 1.0, "scan_max_fps": 1.0, "scan_avg_fps": 1.0}])
        self.assertEqual(rows[0]["mode_name"], "平行线")
        self.assertIsNone(rows[1]["preview"])


class SystemEnvTests(unittest.TestCase):
    def test_system_info_parsing_and_version_completion(self) -> None:
        info = parse_system_info_text(
            "CPU: Intel(R) Core(TM) Ultra 7 255HX\n"
            "CPU Cores: 20\n"
            "Memory: 32175MB\n"
            "GPU: NVIDIA GeForce RTX 5060 Laptop GPU\n"
            "GPU Memory: 7899MB\n"
            "OS: Windows 11  (10.0.22631) 64bit\n"
        )
        self.assertEqual(info["cpu_cores"], "20")
        self.assertEqual(info["memory_mb"], 32175)
        self.assertEqual(info["os_name"], "Windows 11 专业版")
        self.assertEqual(info["os_version_name"], "23H2")
        self.assertEqual(info["os_build"], "22631")

    def test_sample_app_log_versions(self) -> None:
        app_log = _SAMPLE_ROOT / "app.log"
        self.assertEqual(read_software_version(str(app_log)), "1.12.11")
        self.assertEqual(read_wifi_handle_version(str(app_log)), "1.3.1")

    def test_sample_system_info_from_logs(self) -> None:
        info = read_system_info_from_logs(str(_SAMPLE_ROOT))
        self.assertEqual(info.get("memory_mb"), 32175)
        self.assertEqual(info.get("gpu"), "NVIDIA GeForce RTX 5060 Laptop GPU")


if __name__ == "__main__":
    unittest.main()
