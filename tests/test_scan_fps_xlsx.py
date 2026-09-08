from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from engine.scan_fps_xlsx import write_scan_fps_workbook


class ScanFpsWorkbookTests(unittest.TestCase):
    def test_writes_module_connection_params_and_average_fps(self) -> None:
        case = {
            "name": "Pika stream",
            "steps": [
                {
                    "id": "crealityscan.configure_scan_params_pika",
                    "params": {"preset": "wifi-line-laser-point-cloud"},
                },
                {
                    "id": "crealityscan.scan_until_frames_then_stop",
                    "name": "scan 200 frames",
                },
            ],
        }
        step_results = [
            {
                "id": "crealityscan.configure_scan_params_pika",
                "status": "passed",
                "extra": {},
            },
            {
                "id": "crealityscan.scan_until_frames_then_stop",
                "name": "scan 200 frames",
                "status": "passed",
                "extra": {
                    "avg_fps": 119.968,
                    "target_frames": 200,
                    "start_frame": 3,
                    "stop_click_frame": 203,
                    "frame_delta": 200,
                    "scan_elapsed_sec": 1.668,
                    "scan_started_log_at": "2026-08-10T20:50:48.268637",
                    "stop_click_log_at": "2026-08-10T20:50:49.936637",
                },
            },
            {
                "id": "crealityscan.fusion_operation",
                "status": "passed",
                "duration_sec": 21.058,
                "extra": {"operation_elapsed_sec": 20.732},
            },
            {
                "id": "crealityscan.package_operation",
                "status": "passed",
                "duration_sec": 7.828,
                "extra": {"operation_elapsed_sec": 7.515},
            },
            {
                "id": "crealityscan.texture_operation",
                "status": "passed",
                "duration_sec": 14.225,
                "extra": {"operation_elapsed_sec": 13.867},
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = write_scan_fps_workbook(
                case,
                step_results,
                temp_dir,
                device_info={"camera_connection_type": "USB3.0"},
            )

            self.assertIsNotNone(workbook_path)
            workbook = load_workbook(Path(str(workbook_path)), data_only=False)
            worksheet = workbook[workbook.sheetnames[0]]
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(worksheet.cell(2, 3).value, "Pika")
            self.assertEqual(worksheet.cell(2, 4).value, "USB3.0")
            self.assertEqual(worksheet.cell(2, 5).value, "wifi-line-laser-point-cloud")
            self.assertEqual(worksheet.cell(2, 10).value, 119.968)
            self.assertEqual(worksheet.cell(2, 10).number_format, "0.000")
            self.assertEqual(worksheet.cell(2, 11).value, 20.732)
            self.assertEqual(worksheet.cell(2, 12).value, 7.515)
            self.assertEqual(worksheet.cell(2, 13).value, 13.867)
            self.assertEqual(worksheet.tables["ScanFpsTable"].ref, "A1:T2")
            self.assertEqual(worksheet.column_dimensions["E"].width, 24)
            self.assertEqual(worksheet.column_dimensions["T"].width, 32)
            self.assertEqual(worksheet.row_dimensions[2].height, 20)

    def test_skipped_scan_is_not_written(self) -> None:
        case = {
            "name": "skipped",
            "steps": [{"id": "crealityscan.scan_until_frames_then_stop"}],
        }
        step_results = [
            {
                "id": "crealityscan.scan_until_frames_then_stop",
                "status": "skipped",
                "extra": {"avg_fps": 30.0},
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = write_scan_fps_workbook(case, step_results, temp_dir)

            self.assertIsNone(workbook_path)


if __name__ == "__main__":
    unittest.main()
